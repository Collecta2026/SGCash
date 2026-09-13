"""Tests for the cash position strip, completeness, delegation, charts,
uploads and the weekly notification flow.

As in `test_engine.py`, every asserted figure is computed independently
of the code under test.
"""
import os
import sys
from datetime import date, timedelta
from decimal import Decimal

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from models import (db, User, CustomerCollection, SupplierInstalment,
                    BankLoanInstalment, FixedWeeklyCost, BankAccount,
                    NotificationLog, set_setting, ST_APPROVED, ST_DRAFT)
import services as svc
import permissions as perms
import charts
import importer


@pytest.fixture()
def app():
    from app import create_app
    a = create_app({"SQLALCHEMY_DATABASE_URI": "sqlite://", "TESTING": True,
                    "SECRET_KEY": "test", "WTF_CSRF_ENABLED": False})
    with a.app_context():
        db.drop_all()
        db.create_all()
        for k, v in svc.DEFAULTS.items():
            set_setting(k, v)
        set_setting("forecast_start", "2026-01-04")
        set_setting("week_start", "6")
        set_setting("work_week", "sun_thu")
        set_setting("weekend_rule", "next")
        set_setting("fx_rate", "50")
        set_setting("horizon_weeks", "52")
        perms.seed_matrix(force=True)
        yield a
        db.session.remove()


def bank(ccy, amount):
    db.session.add(BankAccount(name=f"{ccy} account", currency=ccy,
                               balance=Decimal(str(amount)), as_at=date(2026, 1, 4)))
    db.session.commit()


def add(model, **kw):
    kw.setdefault("status", ST_APPROVED)
    kw.setdefault("currency", "EGP")
    row = model(**kw)
    db.session.add(row)
    db.session.commit()
    return row


# ============================================================
# Cash position strip
# ============================================================

def test_strip_spans_last_week_this_week_and_ahead(app):
    bank("EGP", 100000)
    s = svc.cash_position_strip(back=1, ahead=3)
    assert len(s["rows"]) == 5
    offsets = [r["offset"] for r in s["rows"]]
    assert offsets == [-1, 0, 1, 2, 3]
    assert s["current"]["offset"] == 0
    assert sum(1 for r in s["rows"] if r["is_past"]) == 1


def test_this_weeks_opening_is_the_bank_position(app):
    bank("EGP", 250000)
    bank("USD", 4000)
    s = svc.cash_position_strip(back=2, ahead=2)
    assert s["current"]["opening"]["EGP"] == Decimal("250000")
    assert s["current"]["opening"]["USD"] == Decimal("4000")


def test_history_is_reconstructed_backwards_consistently(app):
    """Last week's closing must equal this week's opening, and last week's
    opening must follow from its own movements."""
    bank("EGP", 100000)
    weeks = svc.build_weeks(n=2, start=svc.forecast_start() - timedelta(days=7))
    add(CustomerCollection, customer="Paid last week", amount=Decimal("30000"),
        due_date=weeks[0]["start"] + timedelta(days=1))
    add(BankLoanInstalment, lender="Bank", amount=Decimal("12000"),
        due_date=weeks[0]["start"] + timedelta(days=2))
    s = svc.cash_position_strip(back=1, ahead=1)
    last, now = s["rows"][0], s["rows"][1]
    assert last["closing"]["EGP"] == now["opening"]["EGP"] == Decimal("100000")
    # opening = closing - in + out
    assert last["opening"]["EGP"] == Decimal("100000") - Decimal("30000") + Decimal("12000")
    assert last["opening"]["EGP"] == Decimal("82000")


def test_strip_arithmetic_holds_in_every_row(app):
    bank("EGP", 60000)
    bank("USD", 3000)
    add(FixedWeeklyCost, description="Wages", weekday=1, amount=Decimal("4000"))
    add(CustomerCollection, customer="A", amount=Decimal("1000"), currency="USD",
        due_date=svc.forecast_start() + timedelta(days=8))
    s = svc.cash_position_strip(back=2, ahead=4)
    for r in s["rows"]:
        for c in ("EGP", "USD"):
            assert r["closing"][c] == r["opening"][c] + r["inflow"][c] - r["outflow"][c]
    for i in range(1, len(s["rows"])):
        for c in ("EGP", "USD"):
            assert s["rows"][i]["opening"][c] == s["rows"][i - 1]["closing"][c]


def test_gap_is_measured_against_the_buffer_not_zero(app):
    set_setting("min_buffer_EGP", "50000")
    bank("EGP", 60000)
    add(BankLoanInstalment, lender="Bank", amount=Decimal("25000"),
        due_date=svc.forecast_start() + timedelta(days=1))
    s = svc.cash_position_strip(back=0, ahead=1)
    now = s["current"]
    assert now["closing"]["EGP"] == Decimal("35000")
    assert now["gap"]["EGP"] == Decimal("35000") - Decimal("50000")
    assert now["gap"]["EGP"] == Decimal("-15000")


def test_no_gap_when_the_balance_clears_the_buffer(app):
    set_setting("min_buffer_EGP", "10000")
    bank("EGP", 60000)
    s = svc.cash_position_strip(back=0, ahead=1)
    assert s["current"]["gap"]["EGP"] == Decimal("0")
    assert s["total_gap_eqv"] == Decimal("0")


# ============================================================
# Completeness
# ============================================================

def test_completeness_flags_a_week_a_table_has_missed(app):
    weeks = svc.build_weeks(4)
    for i in (0, 1, 3):
        add(CustomerCollection, customer=f"W{i}", amount=Decimal("100"),
            due_date=weeks[i]["start"] + timedelta(days=1))
    comp = svc.completeness(weeks=weeks)
    row = next(r for r in comp["rows"] if r["key"] == "collections")
    assert row["counts"] == [1, 1, 0, 1]
    assert row["gaps"] == [2]
    assert row["populated"] == 3
    assert row["coverage"] == Decimal("75")
    assert any(a["key"] == "gap_missing_weeks" and a["stream"] == "collections"
               for a in comp["alerts"])


def test_a_table_never_used_is_reported_once_not_as_weekly_gaps(app):
    weeks = svc.build_weeks(4)
    comp = svc.completeness(weeks=weeks)
    row = next(r for r in comp["rows"] if r["key"] == "petty_cash")
    assert row["never_used"] is True
    assert row["gaps"] == []
    assert any(a["key"] == "gap_never_used" and a["stream"] == "petty_cash"
               for a in comp["alerts"])


def test_completely_empty_weeks_are_flagged(app):
    weeks = svc.build_weeks(3)
    add(CustomerCollection, customer="Only week 0", amount=Decimal("100"),
        due_date=weeks[0]["start"] + timedelta(days=1))
    comp = svc.completeness(weeks=weeks)
    assert comp["empty_weeks"] == [1, 2]
    assert sum(1 for a in comp["alerts"] if a["key"] == "gap_empty_week") == 2


def test_recurring_costs_count_as_present_in_every_week(app):
    """A fixed weekly cost has one row but fills every week — it must not
    be reported as a gap."""
    weeks = svc.build_weeks(5)
    add(FixedWeeklyCost, description="Wages", weekday=1, amount=Decimal("900"))
    comp = svc.completeness(weeks=weeks)
    row = next(r for r in comp["rows"] if r["key"] == "fixed_weekly")
    assert row["counts"] == [1] * 5
    assert row["gaps"] == []


def test_missing_for_role_only_reports_that_roles_tables(app):
    weeks = svc.build_weeks(3)
    add(BankLoanInstalment, lender="Bank", amount=Decimal("500"),
        due_date=weeks[0]["start"] + timedelta(days=1))
    _comp, missing = svc.missing_for_role("sales_admin", week_index=0, weeks=weeks)
    keys = {m["key"] for m in missing}
    assert "collections" in keys              # sales admin owns collections
    assert "bank_loans" not in keys           # and is not asked about loans


# ============================================================
# Delegation
# ============================================================

def _user(role):
    u = User(username=f"u_{role}", role=role)
    u.set_password("pw123456")
    db.session.add(u)
    db.session.commit()
    return u


@pytest.mark.parametrize("role", ["credit_control", "sales_admin"])
def test_data_entry_roles_reach_their_own_table_and_nothing_else(app, role):
    """Credit Control and Sales Admin key in customer collections. They see
    no cash position, no other table, and no other page."""
    u = _user(role)
    assert perms.can(u, "collections", "view")
    assert perms.can(u, "collections", "enter")
    assert perms.can(u, "collections", "edit")
    for other in ("suppliers", "bank_loans", "cheques", "petty_cash",
                  "fixed_weekly", "fixed_monthly", "refunds", "adhoc_in", "one_off"):
        for action in ("view", "enter", "edit", "delete", "approve"):
            assert not perms.can(u, other, action), f"{role} reached {other}.{action}"


@pytest.mark.parametrize("role", ["credit_control", "sales_admin"])
def test_data_entry_roles_hold_no_core_capability(app, role):
    u = _user(role)
    for cap in perms.CORE_CAP_KEYS:
        assert not perms.has_perm(u, cap), f"{role} holds {cap}"


@pytest.mark.parametrize("role", ["credit_control", "sales_admin"])
def test_data_entry_roles_may_not_approve_anything(app, role):
    u = _user(role)
    assert not perms.can_approve_any(u)


@pytest.mark.parametrize("role", ["credit_control", "sales_admin"])
def test_the_menu_lists_only_the_one_table_a_data_entry_role_owns(app, role):
    u = _user(role)
    assert perms.visible_streams(u) == ["collections"]
    assert perms.landing_stream(u) == "collections"
    assert perms.visible_caps(u, perms.CORE_CAP_KEYS) == []


@pytest.mark.parametrize("role", ["md", "cfo", "finance_manager"])
def test_senior_roles_see_every_table_on_the_menu(app, role):
    u = _user(role)
    assert len(perms.visible_streams(u)) == len(perms.STREAM_ORDER)


def test_the_menu_never_lists_something_the_matrix_refuses(app):
    """Whatever the role, every menu entry it is given must be usable."""
    for role in perms.ROLE_KEYS:
        u = _user(role)
        for stream in perms.visible_streams(u):
            assert perms.can(u, stream, "view"), f"{role} listed {stream} it cannot open"
        for cap in perms.visible_caps(u, perms.CAP_KEYS):
            assert perms.has_perm(u, cap), f"{role} listed {cap} it cannot use"


@pytest.mark.parametrize("role", ["md", "cfo", "finance_manager"])
def test_the_three_senior_roles_view_and_edit_the_data(app, role):
    """Finance Manager, CFO and MD are the seats that work the figures."""
    u = _user(role)
    assert perms.has_perm(u, "dashboard")
    assert perms.has_perm(u, "forecast")
    assert perms.has_perm(u, "reports")
    assert perms.has_perm(u, "banks_view")
    for stream in perms.STREAM_ORDER:
        assert perms.can(u, stream, "view")
        assert perms.can(u, stream, "enter")
        assert perms.can(u, stream, "edit")


def test_the_seven_roles_are_exactly_those_agreed(app):
    assert [r[0] for r in perms.ROLES] == [
        "super_admin", "admin", "md", "cfo", "finance_manager",
        "credit_control", "sales_admin"]
    assert perms.ROLE_LABELS["finance_manager"] == "Finance Manager"
    assert "viewer" not in perms.ROLE_KEYS


def test_only_the_super_administrator_rewrites_the_rules(app):
    sup = _user("super_admin")
    adm = _user("admin")
    assert perms.is_super(sup) and perms.is_admin(sup)
    assert perms.is_admin(adm) and not perms.is_super(adm)
    # the console opens for both, the delegation matrix and settings for one
    assert perms.has_perm(sup, "access_control") and perms.has_perm(sup, "settings")
    assert not perms.has_perm(adm, "access_control")
    assert not perms.has_perm(adm, "settings")
    # and an Administrator still runs the system day to day
    for cap in ("dashboard", "forecast", "reports", "imports", "audit_report"):
        assert perms.has_perm(adm, cap), cap


def test_old_role_names_are_migrated_not_left_stranded(app):
    from models import User as U
    for old in ("account_manager", "accountant", "treasury", "data_entry"):
        u = U(username=f"old_{old}", role=old)
        u.set_password("pw123456")
        db.session.add(u)
    stranded = U(username="old_viewer", role="viewer")
    stranded.set_password("pw123456")
    db.session.add(stranded)
    db.session.commit()
    perms.migrate_roles()
    assert U.query.filter_by(username="old_account_manager").first().role == "finance_manager"
    assert U.query.filter_by(username="old_accountant").first().role == "finance_manager"
    assert U.query.filter_by(username="old_treasury").first().role == "credit_control"
    # a role with no mapping is deactivated rather than silently given access
    left = U.query.filter_by(username="old_viewer").first()
    assert left.active is False


def test_only_the_administrator_reaches_the_admin_console(app):
    for r in ("md", "cfo", "account_manager", "credit_control", "sales_admin", "viewer"):
        u = User(username=f"u_{r}", role=r)
        db.session.add(u)
        db.session.commit()
        assert not perms.is_admin(u), r
    admin = User(username="root", role="admin")
    db.session.add(admin)
    db.session.commit()
    assert perms.is_admin(admin)


def test_md_both_reviews_and_edits(app):
    """The MD is one of the three seats that work the figures, so the role
    carries entry and editing as well as approval."""
    u = User(username="md2", role="md")
    db.session.add(u)
    db.session.commit()
    assert perms.can(u, "collections", "approve")
    assert perms.can(u, "collections", "enter")
    assert perms.can(u, "collections", "edit")
    assert perms.has_perm(u, "reports")
    # deletion stays with the Finance Manager and the administrator
    assert not perms.can(u, "collections", "delete")


def test_data_entry_role_is_refused_every_page_but_its_own(app):
    """The permission model is enforced at the route, not just the menu."""
    client = app.test_client()
    _user("credit_control")
    client.post("/login", data={"username": "u_credit_control",
                                "password": "pw123456"}, follow_redirects=True)
    for path in ("/forecast", "/banks", "/opening", "/trend",
                 "/admin", "/reports?report=cashflow", "/reports?report=audit",
                 "/table/suppliers", "/table/bank_loans", "/table/cheques"):
        assert client.get(path).status_code == 403, path
    assert client.get("/table/collections").status_code == 200
    assert client.get("/table/collections/new").status_code == 200
    # and the front page sends them to their table rather than refusing them
    assert client.get("/").status_code == 302


def test_data_entry_dashboard_shows_no_cash_figures(app):
    client = app.test_client()
    _user("sales_admin")
    client.post("/login", data={"username": "u_sales_admin",
                                "password": "pw123456"}, follow_redirects=True)
    body = client.get("/", follow_redirects=True).data.decode()
    assert "Cash position" not in body
    assert "Closing" not in body
    assert 'class="off"' not in body          # nothing dimmed: it is hidden
    assert "Customer Collections" in body


def test_a_refused_action_still_explains_itself(app):
    """The menu no longer dims anything, but in-page buttons still say why
    they are unavailable — in both languages."""
    u = _user("credit_control")
    reason = perms.dim_reason(u, "banks_view", "en")
    assert "View bank balances" in reason and "Credit Controller" in reason
    assert perms.dim_reason(u, "collections_enter", "en") == ""   # granted: no reason
    assert "صلاحية" in perms.dim_reason(u, "banks_view", "ar")


def test_review_and_request_recipients_are_configurable(app):
    assert perms.get_review_roles() == ["md", "cfo"]
    perms.set_review_roles(["md", "cfo", "finance_manager"])
    assert "finance_manager" in perms.get_review_roles()
    perms.set_review_roles(["nonsense"])
    assert perms.get_review_roles() == []


# ============================================================
# Week sign-off and notifications
# ============================================================

def test_confirming_a_week_records_who_and_when(app):
    wid = svc.current_week_id()
    assert svc.get_signoff(wid) is None
    so = svc.confirm_week(wid, "Hala Adel", "account_manager", "Checked against the bank")
    assert so.status == "confirmed"
    assert so.confirmed_by == "Hala Adel"
    assert so.notes == "Checked against the bank"
    assert svc.get_signoff(wid) is not None
    svc.reopen_week(wid)
    assert svc.get_signoff(wid) is None


def test_weekly_review_is_not_sent_for_an_unconfirmed_week(app):
    import notifications
    result = notifications.send_weekly_review(svc.current_week_id())
    assert result["sent"] == 0
    assert result["reason"] == "week not confirmed"


def test_weekly_review_reports_missing_addresses_rather_than_failing(app):
    import notifications
    bank("EGP", 1000)
    db.session.add(User(username="md", full_name="MD", role="md", email=None))
    db.session.commit()
    svc.confirm_week(svc.current_week_id(), "Hala", "account_manager")
    result = notifications.send_weekly_review(svc.current_week_id())
    assert result["sent"] == 0
    assert any("No email address" in r["message"] for r in result["results"])


def test_input_request_names_only_the_recipients_own_gaps(app):
    import notifications
    weeks = svc.build_weeks(3)
    add(BankLoanInstalment, lender="Bank", amount=Decimal("500"),
        due_date=weeks[0]["start"] + timedelta(days=1))
    user = User(username="sa", full_name="Sales Admin", role="sales_admin",
                email="sa@example.com")
    db.session.add(user)
    db.session.commit()
    _comp, missing = svc.missing_for_role("sales_admin", 0, weeks)
    html = notifications.build_input_request(user, weeks[0], missing)
    assert "Customer Collections" in html
    assert "Bank Loan Instalments" not in html
    assert "Sales Admin" in html


def test_scheduler_does_nothing_without_mail_configured(app):
    import notifications
    result = notifications.run_scheduled()
    assert result["ran"] is False
    assert result["reason"] == "email not configured"


def test_request_due_only_inside_the_notice_window(app):
    import notifications
    set_setting("notify_days_before", "3")
    set_setting("forecast_start", (date.today() + timedelta(days=10)).isoformat())
    assert notifications.due_week_for_requests() is None      # too far out
    set_setting("forecast_start", (date.today() + timedelta(days=2)).isoformat())
    assert notifications.due_week_for_requests() is not None   # inside the window


# ============================================================
# Audit trail
# ============================================================

def test_audit_report_summarises_activity_per_user(app):
    svc.audit("Nadia", "collections_created", 1, "x", stream="collections",
              row_id=1, role="credit_control")
    svc.audit("Nadia", "collections_updated", 1, "y", stream="collections", row_id=1)
    svc.audit("Omar", "collections_deleted", 2, "z", stream="collections", row_id=2)
    rep = svc.audit_report()
    assert rep["count"] == 3
    people = {p["actor"]: p for p in rep["people"]}
    assert people["Nadia"]["created"] == 1
    assert people["Nadia"]["updated"] == 1
    assert people["Nadia"]["total"] == 2
    assert people["Omar"]["deleted"] == 1
    assert rep["by_stream"]["collections"] == 3


def test_audit_report_filters_by_user_and_stream(app):
    svc.audit("Nadia", "collections_created", 1, stream="collections", row_id=1)
    svc.audit("Omar", "suppliers_created", 2, stream="suppliers", row_id=2)
    assert svc.audit_report(user="Nadia")["count"] == 1
    assert svc.audit_report(stream="suppliers")["count"] == 1
    assert svc.audit_report(action="created")["count"] == 2
    assert svc.audit_report(user="Nobody")["count"] == 0


# ============================================================
# Supplier payment instalments
# ============================================================

def test_supplier_payment_types_are_grouped_in_the_planner(app):
    weeks = svc.build_weeks(3)
    add(SupplierInstalment, supplier="Siemens", payment_type="deposit",
        amount=Decimal("5000"), currency="USD",
        due_date=weeks[0]["start"] + timedelta(days=1))
    add(SupplierInstalment, supplier="Siemens", payment_type="instalment",
        amount=Decimal("2000"), currency="USD",
        due_date=weeks[0]["start"] + timedelta(days=2))
    add(SupplierInstalment, supplier="Nile", payment_type="instalment",
        amount=Decimal("40000"), due_date=weeks[1]["start"] + timedelta(days=1))
    plan = svc.weekly_stream("suppliers", weeks)
    assert plan["cells"][0]["groups"]["deposit"]["USD"] == Decimal("5000")
    assert plan["cells"][0]["groups"]["instalment"]["USD"] == Decimal("2000")
    assert plan["cells"][1]["totals"]["EGP"] == Decimal("40000")
    assert plan["group_totals"]["instalment"]["EGP"] == Decimal("40000")


def test_supplier_payments_are_outflows_in_the_forecast(app):
    bank("EGP", 100000)
    add(SupplierInstalment, supplier="Nile", payment_type="final",
        amount=Decimal("25000"), due_date=svc.forecast_start() + timedelta(days=1))
    r = svc.build_forecast()["rows"][0]
    assert r["outflow"]["EGP"] == Decimal("25000")
    assert r["closing"]["EGP"] == Decimal("75000")


# ============================================================
# Uploads
# ============================================================

def test_upload_accepts_a_good_row_and_explains_a_bad_one(app):
    from models import RevenueType
    db.session.add(RevenueType(code="EQ", name_en="Equipment sales"))
    db.session.commit()
    records = [
        {"_row": 2, "customer": "Good Hospital", "collection_type": "Instalment",
         "due_date": "12/11/2026", "currency": "EGP", "amount": "55,000",
         "revenue_type_id": "Equipment sales"},
        {"_row": 3, "customer": "Bad Hospital", "collection_type": "Instalment",
         "due_date": "not a date", "currency": "EGP", "amount": "10"},
        {"_row": 4, "customer": "", "collection_type": "Instalment",
         "due_date": "2026-11-12", "currency": "EGP", "amount": "10"},
    ]
    accepted, rejected = importer.validate("collections", records)
    assert len(accepted) == 1
    assert accepted[0]["customer"] == "Good Hospital"
    assert accepted[0]["amount"] == Decimal("55000.00")
    assert accepted[0]["due_date"] == date(2026, 11, 12)      # day first
    assert accepted[0]["collection_type"] == "instalment"
    assert len(rejected) == 2
    assert any("not a date" in e for e in rejected[0]["errors"])
    assert any("required" in e for e in rejected[1]["errors"])


def test_upload_rejects_a_negative_or_zero_amount(app):
    accepted, rejected = importer.validate("collections", [
        {"_row": 2, "customer": "A", "due_date": "2026-11-12", "amount": "-5"},
        {"_row": 3, "customer": "B", "due_date": "2026-11-12", "amount": "0"},
    ])
    assert accepted == []
    assert len(rejected) == 2


def test_upload_maps_a_category_by_name_and_rejects_an_unknown_one(app):
    from models import CostCategory
    db.session.add(CostCategory(code="PR", name_en="Payroll", name_ar="الرواتب"))
    db.session.commit()
    refs = importer.reference_lookup()
    val, err = importer.coerce(
        {"name": "cost_category_id", "type": "cost_category", "en": "Cost"},
        "Payroll", refs)
    assert err is None and val is not None
    val, err = importer.coerce(
        {"name": "cost_category_id", "type": "cost_category", "en": "Cost"},
        "Nonsense", refs)
    assert val is None and "not a known category" in err


def test_uploaded_rows_arrive_as_drafts(app):
    accepted, _ = importer.validate("collections", [
        {"_row": 2, "customer": "Draft Hospital", "due_date": "2026-11-12",
         "currency": "EGP", "amount": "100"}])
    written = importer.commit("collections", accepted, "Tester")
    assert written == 1
    row = CustomerCollection.query.filter_by(customer="Draft Hospital").first()
    assert row.status == ST_DRAFT
    assert row.created_by == "Tester"


def test_template_has_a_column_for_every_field(app):
    data = importer.build_template("collections")
    assert data[:2] == b"PK"                       # a real xlsx
    from openpyxl import load_workbook
    import io as _io
    wb = load_workbook(_io.BytesIO(data))
    ws = wb["Data"]
    headers = [c.value for c in next(ws.iter_rows(max_row=1))]
    from streams import STREAMS
    assert len(headers) == len(STREAMS["collections"]["fields"])
    assert any(h and h.startswith("Customer no.") for h in headers)
    assert "How to use" in wb.sheetnames


# ============================================================
# Charts
# ============================================================

def test_charts_render_valid_svg(app):
    svg = charts.cash_balance([("W1", 100), ("W2", -50), ("W3", 20)],
                              threshold=10, projected_from=1)
    assert svg.startswith("<svg") and svg.endswith("</svg>")
    assert "polyline" in svg and "Projected" in svg
    assert charts.NEG_FILL in svg                 # the negative week is shaded red


def test_chart_handles_no_data_without_crashing(app):
    for svg in (charts.cash_balance([]), charts.money_in_out([]),
                charts.donut([]), charts.ranked_bars([])):
        assert svg.startswith("<svg")


def test_donut_folds_a_ninth_category_into_other(app):
    slices = [(f"Cat {i}", 10) for i in range(12)]
    svg = charts.donut(slices, max_slices=8)
    assert svg.count("<path") == 8
    assert "Other" in svg


def test_ranked_bars_paint_negatives_red(app):
    svg = charts.ranked_bars([("Surplus", 100), ("Shortfall", -80)])
    assert charts.NEG in svg


def test_chart_escapes_names_rather_than_injecting_markup(app):
    svg = charts.donut([("<script>alert(1)</script>", 10), ("Safe", 5)])
    assert "<script>" not in svg
    assert "&lt;script&gt;" in svg


def test_axis_steps_are_round_numbers(app):
    lo, hi, step = charts._scale(0, 47000)
    assert step in (10000.0, 20000.0, 25000.0, 5000.0)
    assert lo == 0 and hi >= 47000
    lo, hi, _ = charts._scale(-3000, 8000)
    assert lo <= -3000 and hi >= 8000              # the scale always includes zero


# ============================================================
# Available cash: green, amber, red
# ============================================================

def test_cash_health_is_green_with_comfortable_headroom(app):
    bank("EGP", 1000000)
    add(BankLoanInstalment, lender="Bank", amount=Decimal("100000"),
        due_date=svc.forecast_start() + timedelta(days=1))
    h = svc.cash_health(6)
    assert h["available"] == Decimal("1000000")
    assert h["committed_out"] == Decimal("100000")
    assert h["state"] == "green"


def test_cash_health_turns_amber_within_the_headroom_threshold(app):
    """1,100,000 available against 1,000,000 committed is 10% headroom,
    inside the 15% threshold, so amber."""
    set_setting("amber_headroom_pct", "15")
    bank("EGP", 1100000)
    add(BankLoanInstalment, lender="Bank", amount=Decimal("1000000"),
        due_date=svc.forecast_start() + timedelta(days=1))
    h = svc.cash_health(6)
    assert h["headroom"] == Decimal("100000")
    assert h["headroom_pct"] == Decimal("10")
    assert h["state"] == "amber"


def test_cash_health_stays_green_just_outside_the_threshold(app):
    """1,200,000 against 1,000,000 is 20% headroom — clear of 15%."""
    set_setting("amber_headroom_pct", "15")
    bank("EGP", 1200000)
    add(BankLoanInstalment, lender="Bank", amount=Decimal("1000000"),
        due_date=svc.forecast_start() + timedelta(days=1))
    h = svc.cash_health(6)
    assert h["headroom_pct"] == Decimal("20")
    assert h["state"] == "green"


def test_cash_health_is_red_when_a_balance_goes_negative(app):
    bank("EGP", 100000)
    add(BankLoanInstalment, lender="Bank", amount=Decimal("250000"),
        due_date=svc.forecast_start() + timedelta(days=1))
    h = svc.cash_health(6)
    assert h["state"] == "red"
    assert h["shortage_week"] is not None
    assert h["shortage_amount"] < 0


def test_cash_health_is_red_when_cash_does_not_cover_commitments(app):
    """Even without a negative week, cash that cannot meet what is already
    committed is a shortage."""
    bank("EGP", 900000)
    for i in range(6):
        add(CustomerCollection, customer=f"C{i}", amount=Decimal("200000"),
            due_date=svc.forecast_start() + timedelta(days=7 * i + 1))
        add(BankLoanInstalment, lender="Bank", amount=Decimal("200000"),
            due_date=svc.forecast_start() + timedelta(days=7 * i + 2))
    h = svc.cash_health(6)
    assert h["committed_out"] == Decimal("1200000")
    assert h["available"] == Decimal("900000")
    assert h["state"] == "red"


def test_cash_health_threshold_is_configurable(app):
    bank("EGP", 1100000)
    add(BankLoanInstalment, lender="Bank", amount=Decimal("1000000"),
        due_date=svc.forecast_start() + timedelta(days=1))
    set_setting("amber_headroom_pct", "5")
    assert svc.cash_health(6)["state"] == "green"      # 10% clears a 5% threshold
    set_setting("amber_headroom_pct", "25")
    assert svc.cash_health(6)["state"] == "amber"      # and trips a 25% one


def test_cash_health_is_green_when_nothing_is_committed(app):
    bank("EGP", 5000)
    h = svc.cash_health(6)
    assert h["committed_out"] == Decimal("0")
    assert h["headroom_pct"] is None
    assert h["state"] == "green"


def test_cash_health_counts_both_currencies_at_the_rate(app):
    bank("EGP", 100000)
    bank("USD", 10000)                      # 500,000 EGP at the test rate of 50
    h = svc.cash_health(6)
    assert h["available"] == Decimal("600000")


# ============================================================
# Available cash: green, amber, red
# ============================================================

def test_cash_health_is_green_with_comfortable_headroom(app):
    bank("EGP", 1000000)
    add(BankLoanInstalment, lender="Bank", amount=Decimal("100000"),
        due_date=svc.forecast_start() + timedelta(days=1))
    h = svc.cash_health(6)
    assert h["available"] == Decimal("1000000")
    assert h["committed_out"] == Decimal("100000")
    assert h["state"] == "green"


def test_cash_health_turns_amber_within_the_headroom_threshold(app):
    """1,100,000 available against 1,000,000 committed is 10% headroom —
    inside the 15% threshold, so amber."""
    set_setting("amber_headroom_pct", "15")
    bank("EGP", 1100000)
    add(BankLoanInstalment, lender="Bank", amount=Decimal("1000000"),
        due_date=svc.forecast_start() + timedelta(days=1))
    h = svc.cash_health(6)
    assert h["headroom"] == Decimal("100000")
    assert h["headroom_pct"] == Decimal("10")
    assert h["state"] == "amber"


def test_cash_health_stays_green_just_outside_the_threshold(app):
    """1,200,000 against 1,000,000 is 20% headroom — clear of 15%."""
    set_setting("amber_headroom_pct", "15")
    bank("EGP", 1200000)
    add(BankLoanInstalment, lender="Bank", amount=Decimal("1000000"),
        due_date=svc.forecast_start() + timedelta(days=1))
    h = svc.cash_health(6)
    assert h["headroom_pct"] == Decimal("20")
    assert h["state"] == "green"


def test_cash_health_is_red_when_a_balance_goes_negative(app):
    bank("EGP", 100000)
    add(BankLoanInstalment, lender="Bank", amount=Decimal("250000"),
        due_date=svc.forecast_start() + timedelta(days=1))
    h = svc.cash_health(6)
    assert h["state"] == "red"
    assert h["shortage_week"] is not None
    assert h["shortage_amount"] < 0


def test_cash_health_is_red_when_cash_does_not_cover_commitments(app):
    """Even without a negative week, cash that cannot meet what is already
    committed to go out is a shortage."""
    bank("EGP", 900000)
    for i in range(6):
        add(CustomerCollection, customer=f"C{i}", amount=Decimal("200000"),
            due_date=svc.forecast_start() + timedelta(days=7 * i + 1))
        add(BankLoanInstalment, lender="Bank", amount=Decimal("200000"),
            due_date=svc.forecast_start() + timedelta(days=7 * i + 2))
    h = svc.cash_health(6)
    assert h["committed_out"] == Decimal("1200000")
    assert h["available"] == Decimal("900000")
    assert h["state"] == "red"


def test_cash_health_threshold_is_configurable(app):
    bank("EGP", 1100000)
    add(BankLoanInstalment, lender="Bank", amount=Decimal("1000000"),
        due_date=svc.forecast_start() + timedelta(days=1))
    set_setting("amber_headroom_pct", "5")
    assert svc.cash_health(6)["state"] == "green"      # 10% clears a 5% threshold
    set_setting("amber_headroom_pct", "25")
    assert svc.cash_health(6)["state"] == "amber"      # and trips a 25% one


def test_cash_health_is_green_when_nothing_is_committed(app):
    bank("EGP", 5000)
    h = svc.cash_health(6)
    assert h["committed_out"] == Decimal("0")
    assert h["headroom_pct"] is None
    assert h["state"] == "green"


def test_cash_health_counts_both_currencies_at_the_rate(app):
    bank("EGP", 100000)
    bank("USD", 10000)                      # 500,000 EGP at the test rate of 50
    h = svc.cash_health(6)
    assert h["available"] == Decimal("600000")


# ============================================================
# Bank & cash balances on the dashboard
# ============================================================

def test_bank_balances_lists_every_account_with_its_egp_equivalent(app):
    db.session.add_all([
        BankAccount(name="NBE — current account", bank="National Bank of Egypt",
                    currency="EGP", kind="bank", sort=10,
                    balance=Decimal("1450000"), as_at=date(2026, 1, 4)),
        BankAccount(name="NBE — USD account", bank="National Bank of Egypt",
                    currency="USD", kind="bank", sort=20,
                    balance=Decimal("42000"), as_at=date(2026, 1, 2)),
        BankAccount(name="InstaPay", currency="EGP", kind="wallet", sort=50,
                    balance=Decimal("86000"), as_at=date(2026, 1, 4)),
    ])
    db.session.commit()

    b = svc.bank_balances()
    assert [r["account"].name for r in b["rows"]] == [
        "NBE — current account", "NBE — USD account", "InstaPay"]   # by sort
    assert b["totals"]["EGP"] == Decimal("1536000")                 # 1,450,000 + 86,000
    assert b["totals"]["USD"] == Decimal("42000")
    # 1,536,000 + 42,000 x 50
    assert b["eqv"] == Decimal("3636000")
    usd = b["rows"][1]
    assert usd["eqv"] == Decimal("2100000")
    assert b["count"] == 3
    assert b["oldest_as_at"] == date(2026, 1, 2)


def test_bank_balances_excludes_accounts_kept_out_of_the_forecast(app):
    db.session.add_all([
        BankAccount(name="In", currency="EGP", balance=Decimal("100000"),
                    as_at=date(2026, 1, 4), include_in_forecast=True),
        BankAccount(name="Out", currency="EGP", balance=Decimal("900000"),
                    as_at=date(2026, 1, 4), include_in_forecast=False),
    ])
    db.session.commit()

    b = svc.bank_balances()
    assert len(b["rows"]) == 2                       # both are shown
    assert b["totals"]["EGP"] == Decimal("100000")   # only one is counted
    assert b["count"] == 1
    assert [r["account"].name for r in b["excluded"]] == ["Out"]
    # the excluded row still carries its own figures
    assert b["rows"][1]["balance"] == Decimal("900000")
    assert b["rows"][1]["included"] is False


def test_bank_balances_totals_match_the_forecast_opening(app):
    bank("EGP", 100000)
    bank("USD", 10000)
    b = svc.bank_balances()
    assert b["eqv"] == svc.cash_health(6)["available"]


def test_standard_accounts_are_seeded_once_and_balances_are_never_reset(app):
    import seed
    seed.seed_accounts()
    names = [a.name for a in BankAccount.query.order_by(BankAccount.sort).all()]
    assert names == [n for n, _b, _c, _k, _s in seed.STANDARD_ACCOUNTS]
    assert {a.currency for a in BankAccount.query.all()} == {"EGP", "USD"}
    assert [a.kind for a in BankAccount.query.order_by(BankAccount.sort).all()][-1] == "cash"

    nbe = BankAccount.query.filter_by(name="NBE — current account").one()
    nbe.balance = Decimal("777000")
    db.session.commit()

    seed.seed_accounts()                               # running again is a no-op
    assert BankAccount.query.count() == len(seed.STANDARD_ACCOUNTS)
    assert BankAccount.query.filter_by(
        name="NBE — current account").one().balance == Decimal("777000")


def test_the_named_banks_and_wallets_are_all_present(app):
    import seed
    seed.seed_accounts()
    blob = " ".join(a.name + " " + (a.bank or "") for a in BankAccount.query.all())
    for wanted in ("NBE", "National Bank of Egypt", "CIB",
                   "Commercial International Bank", "InstaPay", "Vodafone Cash"):
        assert wanted in blob


# ============================================================
# The menu must be readable
# ============================================================

def _css():
    import pathlib
    return pathlib.Path(__file__).resolve().parent.parent.joinpath(
        "static", "style.css").read_text()


def test_dropdown_entries_win_over_the_white_nav_rule(app):
    """`nav.main a` paints menu links white. The dropdown panel is white, so
    its own rule has to be at least as specific or the entries vanish."""
    css = _css()
    assert "nav.main .dd-menu a" in css, (
        "the dropdown rule is not qualified with nav.main — its links will be "
        "white on a white panel")
    i_nav = css.index("nav.main a, nav.main .dd > span")
    i_menu = css.index("nav.main .dd-menu a")
    assert i_menu > i_nav, "the dropdown rule must come after the white nav rule"


def test_dropdown_opens_on_focus_as_well_as_hover(app):
    css = _css()
    assert ":focus-within .dd-menu" in css, "the menu cannot be opened without a mouse"
